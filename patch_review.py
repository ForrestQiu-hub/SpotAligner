# -*- coding: utf-8 -*-
"""Small corrections to the user's edited literature review.

Works on their file, keeps its formatting: existing paragraphs are rewritten
through their own runs, added paragraphs are copies of a neighbouring paragraph
so that style, font and spacing come along unchanged. Writes a new file.
"""

import copy
import shutil
from docx import Document
from docx.text.paragraph import Paragraph

SRC = r"D:\VST3 Programming\SpotAligner\文献综述-点话筒声像定位与时间差.docx"
DST = r"D:\VST3 Programming\SpotAligner\文献综述-点话筒声像定位与时间差-v2.docx"

shutil.copyfile(SRC, DST)
doc = Document(DST)

log = []


def find(prefix):
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith(prefix):
            return i, p
    return -1, None


def rewrite(par, text):
    """Replace a paragraph's text without touching its formatting."""
    runs = par.runs
    if not runs:
        par.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def insert_after(par, text):
    """A new paragraph that is a copy of this one, so it looks the same."""
    new_p = copy.deepcopy(par._p)
    par._p.addnext(new_p)
    np = Paragraph(new_p, par._parent)
    rewrite(np, text)
    return np


# ---------------------------------------------------------------- correction 1
# section 五: the mono question has two levels, not one
i, p = find("由此可以推得两种做法的差别")
if p is None:
    log.append("MISS correction 1")
else:
    rewrite(p, "由此可以推得两种做法的差别,但必须分成两层来看,两层的结论并不相同。")
    n1 = insert_after(p, (
        "第一层,只看点话筒这一轨。左右两路延时不同,意味着这一轨自身就带有 "
        "Δt = (d_L − d_R)/c 的声道间时间差,切单声道时它自己就有梳状零点,频率为 "
        "f = 1/(2Δt)。例如 Δt = 0.87 ms 时第一个零点落在 575 Hz,正是音乐最敏感的频段。"
        "这一层是确定的:这样一轨,单独按单声道听,确实不干净。"))
    n2 = insert_after(n1, (
        "第二层,把它与主对叠在一起再切单声道。点话筒的零点落在 1/(2Δt),而主对对该声源本应"
        "产生的零点也落在同一个频率上;两者重合时,整体并不新增零点。所以“整体不额外变差”"
        "是有条件的,条件是延时必须以主话筒直达声的差值为参照、并且点话筒本身是一个单声道对象。"
        "若以绝对距离为参照、几何填写有误、或点话筒自身是一对话筒,零点就会错开,整体便真的"
        "多出一层梳状滤波。"))
    n3 = insert_after(n2, (
        "还要说明一点:宿主软件的音量声像旋钮与立体声宽度只改变电平,不产生任何延时,"
        "因此它无法消除梳状,只能通过压低一侧的电平来减轻它。这既解释了“单一延时加声像”"
        "为什么在实际中显得安全,也说明了这种安全来自电平掩蔽,而不是来自结构上的干净。"))
    log.append("OK correction 1 (rewrote 1, added 3)")

# ---------------------------------------------------------------- correction 2
# section 六: the failure modes of a stereo spot pair, stated precisely
i, p = find("工程上的另一个后果")
if p is None:
    log.append("MISS correction 2")
else:
    rewrite(p, (
        "把它与主对叠加再切单声道时,零点是否重合,取决于总时间差是否与主对的一致。"
        "给整对施加同一个延时,内部 τ 完全不动,因此这一做法更稳妥;而给两路分别施加"
        "不同的延时,零点会随着额外延时的大小与方向移动,可能错开主对的零点。"))
    n4 = insert_after(p, (
        "把上面两条合起来,三种失败形态可以说得很具体:当额外延时与 τ 方向相反而大小接近时,"
        "总时间差趋近于零,这对话筒的宽度塌陷成近似单声道;当额外延时更大时,总时间差反号,"
        "像场左右翻转;当总时间差越过约 1 ms 时,低频与高频所依据的定位线索不再一致,"
        "像场开始分裂。这三种情形都不需要主对参与,单看这对点话筒就会出现。"))
    log.append("OK correction 2 (rewrote 1, added 1)")

# ---------------------------------------------------------------- extra sentence
# the one-line summary of the whole argument, at the end of section 六
i, p = find("把上面两条合起来")
if p is not None:
    insert_after(p, (
        "归结成一句:共同延时只是把这一对整体推后,内部关系一点不动,它是对齐;"
        "差动延时必然改写这一对自带的声像,它是声像处理。两者都成立,但不能互相冒充。"))
    log.append("OK closing line added")
else:
    log.append("MISS closing line")

doc.save(DST)
with open(r"D:\VST3 Programming\SpotAligner\patch-log.txt", "w", encoding="utf-8") as f:
    f.write("\n".join(log) + "\n")
print("\n".join(x.encode("ascii", "replace").decode() for x in log))
print("paragraphs:", len(doc.paragraphs), "tables:", len(doc.tables))
